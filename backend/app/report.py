"""
Builds the exportable risk-assessment document (PRD Section 7) from the
draft the user has reviewed and edited client-side, and renders it as
either .docx (python-docx) or .pdf (reportlab) from the same content --
kept as two builders over one input shape, not a docx-to-pdf conversion,
so neither format depends on external tools like LibreOffice.

Nothing here is persisted -- bytes are generated in memory and returned
directly to the export response.
"""
import io

PROTOTYPE_NOTICE = (
    "PROTOTYPE OUTPUT -- not authorised for use on real, current cases "
    "involving identifiable individuals until independently reviewed for "
    "information security and data protection compliance."
)


def _mode_label(mode: str) -> str:
    return "Notifiable Inappropriate Association" if mode == "NIA" else "Business Interest"


def build_docx(draft: dict) -> bytes:
    import docx
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    d = docx.Document()

    notice = d.add_paragraph()
    run = notice.add_run(PROTOTYPE_NOTICE)
    run.bold = True
    run.font.color.rgb = RGBColor(0xB0, 0x00, 0x00)
    run.font.size = Pt(9)
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER

    d.add_heading(f"Risk Assessment -- {_mode_label(draft['mode'])}", level=1)

    header = d.add_table(rows=0, cols=2)
    for label, value in [
        ("Case reference", draft.get("case_reference") or "(not provided)"),
        ("Date", draft.get("date") or ""),
        ("Mode", draft["mode"]),
        ("Preparer", draft.get("preparer") or ""),
    ]:
        row = header.add_row().cells
        row[0].text = label
        row[1].text = value

    d.add_heading("Declared facts summary", level=2)
    facts_table = d.add_table(rows=0, cols=2)
    for item in draft.get("facts_summary", []):
        row = facts_table.add_row().cells
        row[0].text = item.get("label", "")
        row[1].text = item.get("value") or "(not stated)"

    d.add_heading("Risk grading", level=2)
    grade_p = d.add_paragraph()
    grade_run = grade_p.add_run(draft.get("grade") or "(not selected)")
    grade_run.bold = True
    grade_run.font.size = Pt(16)
    if draft["mode"] == "NIA" and draft.get("manageability"):
        m_p = d.add_paragraph()
        m_run = m_p.add_run(draft["manageability"])
        m_run.bold = True
        if draft["manageability"] == "Escalate for manageability review":
            m_run.font.color.rgb = RGBColor(0xB0, 0x00, 0x00)

    d.add_heading("Rationale", level=2)
    d.add_paragraph(draft.get("rationale_text") or "")

    d.add_heading("Recommended conditions/mitigations", level=2)
    d.add_paragraph(draft.get("conditions_text") or "(none)")

    d.add_heading("Recommended review date", level=2)
    d.add_paragraph(draft.get("review_date") or "(not applicable)")

    d.add_heading("Policy basis", level=2)
    d.add_paragraph(
        "This assessment is informed by the following national frameworks: "
        + ", ".join(draft.get("policy_basis", []))
        + ". It does not reference any force-specific internal policy."
    )

    d.add_heading("Sign-off", level=2)
    sign_table = d.add_table(rows=0, cols=2)
    for label, value in [
        ("Preparer name", draft.get("preparer_name") or ""),
        ("Preparer date", draft.get("preparer_date") or ""),
        ("Reviewing officer name", ""),
        ("Reviewing officer date", ""),
    ]:
        row = sign_table.add_row().cells
        row[0].text = label
        row[1].text = value

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def build_pdf(draft: dict) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    def safe_html(text: str) -> str:
        """Escape free text before it reaches Paragraph -- Paragraph parses
        a mini-XML subset, so an officer typing an ordinary '&' or '<' in
        the rationale/conditions/review-date fields would otherwise break
        the export outright. Escape first, then add <br/> line breaks."""
        return escape(str(text or "")).replace("\n", "<br/>")

    styles = getSampleStyleSheet()
    notice_style = ParagraphStyle("notice", parent=styles["Normal"], textColor=HexColor("#b00000"), alignment=1, fontSize=9)
    h1 = styles["Heading1"]
    h2 = styles["Heading2"]
    body = styles["BodyText"]
    grade_style = ParagraphStyle("grade", parent=styles["Normal"], fontSize=20, leading=24)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=20 * mm, bottomMargin=20 * mm)
    story = [
        Paragraph(PROTOTYPE_NOTICE, notice_style),
        Spacer(1, 10),
        Paragraph(f"Risk Assessment -- {_mode_label(draft['mode'])}", h1),
    ]

    header_rows = [
        ["Case reference", draft.get("case_reference") or "(not provided)"],
        ["Date", draft.get("date") or ""],
        ["Mode", draft["mode"]],
        ["Preparer", draft.get("preparer") or ""],
    ]
    story.append(_table(header_rows))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Declared facts summary", h2))
    fact_rows = [[item.get("label", ""), item.get("value") or "(not stated)"] for item in draft.get("facts_summary", [])]
    story.append(_table(fact_rows) if fact_rows else Paragraph("(none)", body))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Risk grading", h2))
    story.append(Paragraph(draft.get("grade") or "(not selected)", grade_style))
    if draft["mode"] == "NIA" and draft.get("manageability"):
        colour = "#b00000" if draft["manageability"] == "Escalate for manageability review" else "#000000"
        story.append(Paragraph(f'<font color="{colour}"><b>{draft["manageability"]}</b></font>', body))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Rationale", h2))
    story.append(Paragraph(safe_html(draft.get("rationale_text")), body))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Recommended conditions/mitigations", h2))
    story.append(Paragraph(safe_html(draft.get("conditions_text")) or "(none)", body))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Recommended review date", h2))
    story.append(Paragraph(safe_html(draft.get("review_date")) or "(not applicable)", body))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Policy basis", h2))
    story.append(Paragraph(
        "This assessment is informed by the following national frameworks: "
        + ", ".join(draft.get("policy_basis", []))
        + ". It does not reference any force-specific internal policy.", body,
    ))
    story.append(Spacer(1, 12))

    story.append(Paragraph("Sign-off", h2))
    sign_rows = [
        ["Preparer name", draft.get("preparer_name") or ""],
        ["Preparer date", draft.get("preparer_date") or ""],
        ["Reviewing officer name", ""],
        ["Reviewing officer date", ""],
    ]
    story.append(_table(sign_rows))

    doc.build(story)
    return buf.getvalue()


def _table(rows):
    """rows: list of [label, value] plain strings. Cell text is wrapped in
    Paragraph flowables -- a plain string in a reportlab Table cell is
    drawn at its natural width and does not wrap, so any value longer than
    its column just overflows and visually bleeds into the next cell/row
    rather than wrapping onto a second line."""
    from xml.sax.saxutils import escape

    from reportlab.lib.colors import HexColor
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, Table, TableStyle

    base = getSampleStyleSheet()["Normal"]
    label_style = ParagraphStyle("cell_label", parent=base, fontSize=9, leading=12, fontName="Helvetica-Bold")
    value_style = ParagraphStyle("cell_value", parent=base, fontSize=9, leading=12)

    wrapped = [
        [Paragraph(escape(str(label)), label_style), Paragraph(escape(str(value)), value_style)]
        for label, value in rows
    ]
    # A4 usable width is ~451pt with default L/R margins -- stay safely under it.
    t = Table(wrapped, colWidths=[130, 300])
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, HexColor("#cccccc")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 0), (0, -1), HexColor("#f2f2f2")),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    return t
